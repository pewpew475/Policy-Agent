"""
Document Processing Service
Handles OCR, PDF processing, and text extraction using tesseract and PyPDF2
"""

import os
import uuid
import asyncio
from pathlib import Path
from typing import Dict, List, Optional, Tuple
import aiofiles
from PIL import Image
import pytesseract
import PyPDF2
# from pdf2image import convert_from_path  # Optional dependency
from docx import Document as DocxDocument
from loguru import logger

from utils.config import settings
from models.schemas import DocumentType, ProcessingStatus

class DocumentProcessor:
    """Document processing service with OCR and PDF capabilities"""
    
    def __init__(self):
        self.upload_path = Path(settings.UPLOAD_PATH)
        self.processed_path = Path(settings.PROCESSED_DOCUMENTS_PATH)
        self.upload_path.mkdir(exist_ok=True)
        self.processed_path.mkdir(exist_ok=True)
        
        # Configure tesseract path if provided
        if settings.TESSERACT_PATH:
            pytesseract.pytesseract.tesseract_cmd = settings.TESSERACT_PATH
    
    async def process_upload(
        self, 
        file_content: bytes, 
        filename: str, 
        content_type: str
    ) -> Dict:
        """Process uploaded file and extract text"""
        
        # Generate unique document ID
        document_id = str(uuid.uuid4())
        file_extension = Path(filename).suffix.lower()
        
        # Determine document type
        document_type = self._determine_document_type(file_extension, content_type)
        
        # Save uploaded file
        safe_filename = f"{document_id}_{filename}"
        file_path = self.upload_path / safe_filename
        
        async with aiofiles.open(file_path, 'wb') as f:
            await f.write(file_content)
        
        logger.info(f"Saved uploaded file: {safe_filename}")
        
        # Process document based on type
        try:
            processing_result = await self._process_document(
                file_path, document_type, document_id
            )
            
            return {
                "document_id": document_id,
                "filename": safe_filename,
                "original_filename": filename,
                "file_path": str(file_path),
                "file_size": len(file_content),
                "document_type": document_type,
                "processing_status": ProcessingStatus.COMPLETED,
                **processing_result
            }
            
        except Exception as e:
            logger.error(f"Error processing document {document_id}: {str(e)}")
            return {
                "document_id": document_id,
                "filename": safe_filename,
                "original_filename": filename,
                "file_path": str(file_path),
                "file_size": len(file_content),
                "document_type": document_type,
                "processing_status": ProcessingStatus.FAILED,
                "error": str(e)
            }
    
    def _determine_document_type(self, file_extension: str, content_type: str) -> DocumentType:
        """Determine document type from extension and content type"""
        
        if file_extension == '.pdf':
            return DocumentType.PDF
        elif file_extension in ['.jpg', '.jpeg', '.png', '.tiff', '.bmp', '.gif']:
            return DocumentType.IMAGE
        elif file_extension in ['.docx', '.doc']:
            return DocumentType.DOCX
        elif file_extension == '.txt':
            return DocumentType.TXT
        else:
            # Fallback to content type
            if 'pdf' in content_type:
                return DocumentType.PDF
            elif 'image' in content_type:
                return DocumentType.IMAGE
            else:
                return DocumentType.TXT
    
    async def _process_document(
        self, 
        file_path: Path, 
        document_type: DocumentType, 
        document_id: str
    ) -> Dict:
        """Process document based on its type"""
        
        if document_type == DocumentType.PDF:
            return await self._process_pdf(file_path, document_id)
        elif document_type == DocumentType.IMAGE:
            return await self._process_image(file_path, document_id)
        elif document_type == DocumentType.DOCX:
            return await self._process_docx(file_path, document_id)
        elif document_type == DocumentType.TXT:
            return await self._process_text(file_path, document_id)
        else:
            raise ValueError(f"Unsupported document type: {document_type}")
    
    async def _process_pdf(self, file_path: Path, document_id: str) -> Dict:
        """Process PDF file with text extraction and OCR fallback"""
        
        extracted_text = ""
        page_count = 0
        
        try:
            # First, try to extract text directly from PDF
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                page_count = len(pdf_reader.pages)
                
                # Check if we exceed page limit
                if page_count > 100:
                    logger.warning(f"PDF has {page_count} pages, processing first 100")
                    page_count = 100
                
                for page_num in range(min(page_count, 100)):
                    page = pdf_reader.pages[page_num]
                    text = page.extract_text()
                    extracted_text += text + "\n"
            
            # If no text extracted or very little text, use OCR
            if len(extracted_text.strip()) < 100:
                logger.info(f"PDF {document_id} has little text, using OCR")
                extracted_text = await self._ocr_pdf(file_path, document_id)
            
        except Exception as e:
            logger.error(f"Error processing PDF {document_id}: {str(e)}")
            # Fallback to OCR
            extracted_text = await self._ocr_pdf(file_path, document_id)
        
        return {
            "extracted_text": extracted_text.strip(),
            "page_count": page_count,
            "confidence_score": 0.9 if len(extracted_text.strip()) > 100 else 0.7
        }
    
    async def _ocr_pdf(self, file_path: Path, document_id: str) -> str:
        """Perform OCR on PDF by converting to images"""

        extracted_text = "OCR processing not available - pdf2image not installed"

        try:
            # Try to import pdf2image dynamically
            from pdf2image import convert_from_path

            # Convert PDF to images
            images = convert_from_path(
                file_path,
                dpi=200,
                first_page=1,
                last_page=100  # Limit to 100 pages
            )

            # Process each page with OCR
            extracted_text = ""
            for i, image in enumerate(images):
                logger.info(f"OCR processing page {i+1} of {document_id}")

                # Perform OCR
                text = pytesseract.image_to_string(
                    image,
                    config='--psm 6 --oem 3'
                )
                extracted_text += text + "\n"

        except ImportError:
            logger.warning("pdf2image not available, skipping OCR for PDF")
            extracted_text = "OCR processing not available - pdf2image not installed"
        except Exception as e:
            logger.error(f"OCR failed for PDF {document_id}: {str(e)}")
            extracted_text = "OCR processing failed"

        return extracted_text
    
    async def _process_image(self, file_path: Path, document_id: str) -> Dict:
        """Process image file with OCR"""
        
        try:
            # Open and process image
            image = Image.open(file_path)
            
            # Perform OCR
            extracted_text = pytesseract.image_to_string(
                image, 
                config='--psm 6 --oem 3'
            )
            
            return {
                "extracted_text": extracted_text.strip(),
                "page_count": 1,
                "confidence_score": 0.8
            }
            
        except Exception as e:
            logger.error(f"Error processing image {document_id}: {str(e)}")
            return {
                "extracted_text": "Image processing failed",
                "page_count": 1,
                "confidence_score": 0.0
            }
    
    async def _process_docx(self, file_path: Path, document_id: str) -> Dict:
        """Process DOCX file"""
        
        try:
            doc = DocxDocument(file_path)
            extracted_text = ""
            
            for paragraph in doc.paragraphs:
                extracted_text += paragraph.text + "\n"
            
            return {
                "extracted_text": extracted_text.strip(),
                "page_count": len(doc.paragraphs) // 20,  # Rough estimate
                "confidence_score": 0.95
            }
            
        except Exception as e:
            logger.error(f"Error processing DOCX {document_id}: {str(e)}")
            return {
                "extracted_text": "DOCX processing failed",
                "page_count": 1,
                "confidence_score": 0.0
            }
    
    async def _process_text(self, file_path: Path, document_id: str) -> Dict:
        """Process plain text file"""
        
        try:
            async with aiofiles.open(file_path, 'r', encoding='utf-8') as f:
                extracted_text = await f.read()
            
            return {
                "extracted_text": extracted_text.strip(),
                "page_count": len(extracted_text) // 2000,  # Rough estimate
                "confidence_score": 1.0
            }
            
        except Exception as e:
            logger.error(f"Error processing text file {document_id}: {str(e)}")
            return {
                "extracted_text": "Text file processing failed",
                "page_count": 1,
                "confidence_score": 0.0
            }
